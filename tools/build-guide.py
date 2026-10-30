#!/usr/bin/env python3
"""Build a game's Players' Guide from <game>/players-guide.md (and its GM's Guide from <game>/gm-guide.md, if it has one).

Writes the guide into the creator's "Players' Guide" tab (between the GUIDE markers
in <game>/index.html) and, with --docx, a Word version.

  python3 tools/build-guide.py [--game uss-pelion|grimm-west] [--docx [--reference-doc original.docx]]
"""
import argparse, json, pathlib, re, subprocess, sys
from stamp import stamp
import gm_guides

ROOT = pathlib.Path(__file__).resolve().parent.parent
# Per game: Word file name, heading color, table header fill.
GAMES = {
    "uss-pelion": ("USS_Pelion_Players_Guide.docx", "1B2A4A", "E9EEF6"),
    "grimm-west": ("Grimm_Players_Guide.docx", "4A2A12", "F1E6D2"),
}
SRC = APP = DOCX = None
NAVY, GOLD_FILL = "1B2A4A", "E9EEF6"
START, END = "<!-- GUIDE:START -->", "<!-- GUIDE:END -->"


def pandoc(*args, src=None):
    return subprocess.run(["pandoc", str(src or SRC), *args], check=True, capture_output=True, text=True).stdout


def front(path, key):
    m = re.search(r"^" + key + r': "(.*)"$', path.read_text().split("\n---", 1)[0], flags=re.M)
    return m.group(1) if m else ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--game", choices=sorted(GAMES), default="uss-pelion", help="which game folder to build")
    ap.add_argument("--docx", action="store_true", help="also write the Word version")
    ap.add_argument("--reference-doc", help="a .docx whose styles the Word version should use")
    a = ap.parse_args()
    global SRC, APP, DOCX, NAVY, GOLD_FILL
    docx_name, NAVY, GOLD_FILL = GAMES[a.game]
    SRC, APP, DOCX = ROOT / a.game / "players-guide.md", ROOT / a.game / "index.html", ROOT / a.game / docx_name

    html = pandoc("-f", "markdown", "-t", "html", "--id-prefix=g-", "--wrap=none")
    # A game with a GM's Guide (<game>/gm-guide.md): both books in the Guide tab, one shown at a time.
    GM = ROOT / a.game / "gm-guide.md"
    # The GM's Guide: in the page, or from the table server for registered GMs only (tools/gm_guides.py, GM_GUIDES_OPEN).
    if GM.exists():
        gm_html = pandoc("-f", "markdown", "-t", "html", "--id-prefix=gm-", "--wrap=none", src=GM)
        gm_guides.write_game(ROOT / a.game, front(GM, "title"), front(GM, "subtitle"), gm_html)
        html = '<div class="guide-book" id="guide-players">\n' + html + "</div>\n" + gm_guides.book(front(GM, "title"), front(GM, "subtitle"), gm_html)
    app = APP.read_text()
    if START not in app or END not in app:
        sys.exit(f"GUIDE markers missing from {APP}")
    app = re.sub(re.escape(START) + r".*?" + re.escape(END), lambda _: f"{START}\n{html}{END}", app, flags=re.S)
    # A page with BATTLEMAP markers gets the shared battle map module (legend-dynamics/engine/battlemap.js).
    if "/* BATTLEMAP:START */" in app:
        bm = (ROOT / "legend-dynamics" / "engine" / "battlemap.js").read_text().replace("</", "<\\/")
        app = re.sub(r"/\* BATTLEMAP:START \*/.*?/\* BATTLEMAP:END \*/", lambda _: "/* BATTLEMAP:START */\n" + bm + "/* BATTLEMAP:END */", app, flags=re.S)
    # A page with ENCOUNTER markers gets the Encounter panel: the shared modules (legend-dynamics/engine) and the
    # game's own setting data (<game>/encounter/setting.json: weapons, threats, gear, combat rules, terms).
    if "/* ENCOUNTER:START */" in app:
        E = ROOT / "legend-dynamics" / "engine"
        setting = (ROOT / a.game / "encounter" / "setting.json").read_text()
        js = "\n".join(["const ENC_SETTING = " + json.dumps(json.loads(setting), ensure_ascii=False, separators=(",", ":")) + ";",
            inline_module(E / "rules" / "core.js", "LD"), inline_module(E / "rules" / "experiments.js", "LDX"),
            inline_module(E / "lab" / "encounter.js", "LAB_ENCOUNTER"), inline_module(E / "lab" / "stats.js", "LAB_STATS"), inline_module(E / "lab" / "ui.js", "LAB_UI")]).replace("</", "<\\/") + "\n"
        app = re.sub(r"/\* ENCOUNTER:START \*/.*?/\* ENCOUNTER:END \*/", lambda _: "/* ENCOUNTER:START */\n" + js + "/* ENCOUNTER:END */", app, flags=re.S)
        app = re.sub(r"<!-- ENCOUNTER-PANEL:START -->.*?<!-- ENCOUNTER-PANEL:END -->", lambda _: "<!-- ENCOUNTER-PANEL:START -->\n" + (E / "lab" / "panel.html").read_text() + "<!-- ENCOUNTER-PANEL:END -->", app, flags=re.S)
    # A page with TABLECLIENT markers gets the shared table's page side (legend-dynamics/engine/table-client.js) and the
    # table server's address (legend-dynamics/table-server/server-url.txt, once it's deployed; otherwise players type it in).
    if "/* TABLECLIENT:START */" in app:
        tc = (ROOT / "legend-dynamics" / "engine" / "table-client.js").read_text().replace("</", "<\\/")
        app = re.sub(r"/\* TABLECLIENT:START \*/.*?/\* TABLECLIENT:END \*/", lambda _: "/* TABLECLIENT:START */\n" + tc + "/* TABLECLIENT:END */", app, flags=re.S)
        f = ROOT / "legend-dynamics" / "table-server" / "server-url.txt"
        url = f.read_text().strip() if f.exists() else ""
        url = url if re.fullmatch(r"https://[\w.-]+(:\d+)?", url) else ""
        app = re.sub(r'const TABLE_SERVER = "[^"]*";', lambda _: f'const TABLE_SERVER = "{url}";', app, count=1)
    # A page with GMLOCK markers locks its GM view and GM's Guide to registered GMs (legend-dynamics/engine/gm-lock.js).
    if "/* GMLOCK:START */" in app:
        gl = (ROOT / "legend-dynamics" / "engine" / "gm-lock.js").read_text().replace("</", "<\\/")
        app = re.sub(r"/\* GMLOCK:START \*/.*?/\* GMLOCK:END \*/", lambda _: "/* GMLOCK:START */\n" + gl + "/* GMLOCK:END */", app, flags=re.S)
    # A page with HOMEBREW markers gets the GM's homebrew threats and weapons (legend-dynamics/engine/homebrew.js, or the
    # game's own copy where the branch has no engine).
    if "/* HOMEBREW:START */" in app:
        src = ROOT / a.game / "homebrew.js"
        src = src if src.exists() else ROOT / "legend-dynamics" / "engine" / "homebrew.js"
        hb = src.read_text().replace("</", "<\\/")
        app = re.sub(r"/\* HOMEBREW:START \*/.*?/\* HOMEBREW:END \*/", lambda _: "/* HOMEBREW:START */\n" + hb + "/* HOMEBREW:END */", app, flags=re.S)
    # A page with PLAIN markers gets the plain hobby look (<game>/plain-theme.css, or legend-dynamics/engine/plain-theme.css),
    # with the game's own palette after it in the page.
    if "/* PLAIN:START */" in app:
        src = ROOT / a.game / "plain-theme.css"
        src = src if src.exists() else ROOT / "legend-dynamics" / "engine" / "plain-theme.css"
        pl = src.read_text()
        app = re.sub(r"/\* PLAIN:START \*/.*?/\* PLAIN:END \*/", lambda _: "/* PLAIN:START */\n" + pl + "/* PLAIN:END */", app, flags=re.S)
    # A page with CAMPAIGN markers gets the campaign builder (<game>/campaign.js, or legend-dynamics/engine/campaign.js).
    if "/* CAMPAIGN:START */" in app:
        src = ROOT / a.game / "campaign.js"
        src = src if src.exists() else ROOT / "legend-dynamics" / "engine" / "campaign.js"
        cp = src.read_text().replace("</", "<\\/")
        app = re.sub(r"/\* CAMPAIGN:START \*/.*?/\* CAMPAIGN:END \*/", lambda _: "/* CAMPAIGN:START */\n" + cp + "/* CAMPAIGN:END */", app, flags=re.S)
    # A page with PLAYTEST markers gets the playtest kit: legend-dynamics/engine/playtest.js and <game>/playtest.json.
    if "/* PLAYTEST:START */" in app:
        kit = (ROOT / a.game / "playtest.json").read_text() if (ROOT / a.game / "playtest.json").exists() else "null"
        pt = (ROOT / "legend-dynamics" / "engine" / "playtest.js").read_text() + "const PLAYTEST_KIT = " + kit.strip() + ";\n"
        app = re.sub(r"/\* PLAYTEST:START \*/.*?/\* PLAYTEST:END \*/", lambda _: "/* PLAYTEST:START */\n" + pt.replace("</", "<\\/") + "/* PLAYTEST:END */", app, flags=re.S)
    app = stamp(app, APP)  # the version footer: a fingerprint of the page, and the day it last changed
    APP.write_text(app)
    print(f"updated {APP.relative_to(ROOT)}")
    gm_guides.bundle()

    if a.docx:
        extra = [f"--reference-doc={a.reference_doc}"] if a.reference_doc else []
        subprocess.run(["pandoc", str(SRC), "-o", str(DOCX), *extra], check=True)
        finish_docx(DOCX)
        print(f"wrote {DOCX.relative_to(ROOT)}")
        if GM.exists():
            gm_docx = DOCX.with_name(DOCX.name.replace("Players_Guide", "GM_Guide"))
            subprocess.run(["pandoc", str(GM), "-o", str(gm_docx), *extra], check=True)
            finish_docx(gm_docx)
            print(f"wrote {gm_docx.relative_to(ROOT)}")


def inline_module(path, name):
    """An ES module as a plain script constant (as tools/build-legend.py does for the creators)."""
    src = path.read_text().replace("// @ts-check\n", "")
    names = re.findall(r"^export (?:async )?(?:function|const|let) (\w+)", src, flags=re.M)
    for line in re.findall(r"^import .*$", src, flags=re.M):
        src = src.replace(line + "\n", "")
    src = re.sub(r"^export (async )?(function|const|let) ", r"\1\2 ", src, flags=re.M)
    return f"const {name} = (() => {{\n{src}\nreturn {{ {', '.join(names)} }};\n}})();"


def finish_docx(path):
    """Give pandoc's plain Word output readable tables and a consistent typeface (needs python-docx)."""
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError:
        print("python-docx not installed; leaving the Word file unstyled")
        return
    doc = Document(str(path))
    def style(name):
        try:
            return doc.styles[name]
        except KeyError:
            return None
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if (st := style(name)) is not None:
            st.font.name, st.font.size = "Calibri", Pt(10.5)
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        if (st := style(name)) is not None:
            st.font.name = "Calibri"
            st.font.color.rgb = RGBColor.from_string(NAVY)
    for table in doc.tables:
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "B8C2D3")
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                if r == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), GOLD_FILL)
                    cell._tc.get_or_add_tcPr().append(shd)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9.5)
                        if r == 0 or c == 0:
                            run.font.bold = True
    doc.save(str(path))


if __name__ == "__main__":
    main()
